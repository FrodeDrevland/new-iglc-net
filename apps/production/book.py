"""The full proceedings (stage 2): one PDF with cover, front matter, the papers and an author index.

Made in the style of the IGLC 32 proceedings:

    front cover                      uploaded (organisers' design, A4)
    colophon                         generated (editors, copyright, ISSN/ISBN)
    title page              i        generated
    conference organisation ii ...   uploaded, from the IGLC template
    foreword                         uploaded, from the template (prefilled with the tables)
    list of reviewers                uploaded, from the template
    other front matter               uploaded, from the template
    table of contents                generated (by track, with track chairs; links to the papers)
    the papers              1 ...    the published PDFs, unchanged
    author index                     generated
    back cover                       uploaded

The uploaded parts leave their footers empty: the system prints the page numbers (roman in the
front matter). Blank pages are added so that the title page and page 1 are right-hand pages
and the back cover is the last left-hand page.
"""

from __future__ import annotations

import io
import re
import unicodedata
import urllib.request
from collections import Counter
from dataclasses import dataclass

from django.core.files.base import ContentFile
from django.core.files.storage import default_storage

from .models import BookPart, Production, Submission, private_storage

PAGE = (595.28, 841.89)  # A4
MARGIN = 70.9            # 2.5 cm, as the papers


class BookError(Exception):
    pass


# ---------------------------------------------------------------- data

def ordinal(n: int) -> str:
    suffix = "th" if 10 <= n % 100 <= 20 else {1: "st", 2: "nd", 3: "rd"}.get(n % 10, "th")
    return f"{n}{suffix}"


def roman(n: int) -> str:
    out = ""
    for value, letters in ((1000, "m"), (900, "cm"), (500, "d"), (400, "cd"), (100, "c"), (90, "xc"),
                           (50, "l"), (40, "xl"), (10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i")):
        while n >= value:
            out, n = out + letters, n - value
    return out


def editors_of(production: Production) -> list[str]:
    return [str(e).strip(" ,;") for e in production.conference.editors.all()]


def names_joined(names: list[str]) -> str:
    return names[0] if len(names) == 1 else ", ".join(names[:-1]) + " and " + names[-1] if names else ""


def place_of(conference) -> str:
    """'Oslo, Norway'; 'Singapore' once."""
    if conference.city == conference.country:
        return conference.city
    return ", ".join(filter(None, [conference.city, conference.country]))


def date_line(conference) -> str:
    from .running import _date_range

    return _date_range(conference.start_date, conference.end_date)


def country_of(affiliation_text: str) -> str:
    """The country of an affiliation: its last part if that is a country, otherwise the last
    country named in it (affiliations are written freely); '' if none is found."""
    from .countries import ALIASES, COUNTRIES
    from .docx_reader import parse_affiliation

    known = {c.lower(): c for c in COUNTRIES} | ALIASES
    rest = parse_affiliation(affiliation_text or "")["affiliation"]
    parts = [re.sub(r"\b\d[\d -]*\b", "", p).strip(" .;()") for p in re.split(r"[,;\n]| - ", rest)]
    for part in reversed([p for p in parts if p]):
        if part.lower() in known:
            return known[part.lower()]
    text = rest.lower()
    found = [(text.rfind(name), country) for name, country in known.items()
             if len(name) > 3 and name not in ("georgia", "jordan") and re.search(rf"\b{re.escape(name)}\b", text)]
    return max(found)[1] if found else ""


def placed_papers(production: Production):
    """[(track or None, [Submission])] in book order, by first page."""
    from .arrange import ordered

    return [(track, sorted(papers, key=lambda s: (s.paper.first_page if s.paper_id else 0)))
            for track, papers in ordered(production)
            if any(s.paper_id for s in papers)]


def statistics(production: Production) -> dict:
    papers = [s.paper for s in production.submissions.filter(paper__isnull=False)
              .select_related("paper").prefetch_related("paper__authors")]
    countries = Counter()
    for paper in papers:
        first = paper.authors.first()
        countries[(country_of(first.title_and_contact) if first else "") or "Unknown"] += 1
    tracks = [(track.title if track else "No track", len([s for s in subs if s.paper_id]))
              for track, subs in placed_papers(production)]
    return {"papers": len(papers), "countries": sorted(countries.items(), key=lambda kv: (-kv[1], kv[0])),
            "tracks": tracks}


# ---------------------------------------------------------------- templates (Word)

TEMPLATE_NOTE = ("[IGLC template for the full proceedings. Keep the styles and leave the header and footer "
                 "empty: the page numbers are added when the proceedings are made. Save as PDF from Word "
                 "(A4) and upload it on the production's Full proceedings page. Delete this note.]")


def _document(heading: str):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()

    def times(r_pr):
        fonts = r_pr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            r_pr.insert(0, fonts)
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            fonts.attrib.pop(qn(attr), None)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            fonts.set(qn(attr), "Times New Roman")

    defaults = doc.styles.element.find(qn("w:docDefaults"))
    if defaults is not None:
        r_pr = defaults.find(qn("w:rPrDefault") + "/" + qn("w:rPr"))
        if r_pr is not None:
            times(r_pr)
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(2.5))
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name, size in (("Heading 1", 12), ("Heading 2", 10)):
        style = doc.styles[name]
        style.font.name, style.font.size, style.font.bold = "Times New Roman", Pt(size), True
        style.font.all_caps, style.font.color.rgb = True, None
        style.paragraph_format.space_before, style.paragraph_format.space_after = Pt(12), Pt(6)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    caption = doc.styles["Caption"]
    caption.font.name, caption.font.size, caption.font.italic, caption.font.bold = "Times New Roman", Pt(10), True, False
    caption.font.color.rgb = None
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for fonts in doc.styles.element.iter(qn("w:rFonts")):  # every style: no theme fonts (Calibri, Cambria)
        times(fonts.getparent())
    note = doc.add_paragraph(TEMPLATE_NOTE)
    note.runs[0].italic = True
    doc.add_heading(heading, level=1)
    return doc


def _table(doc, caption: str, headings: list[str], rows: list[list], total: list | None = None):
    from docx.shared import Pt

    doc.add_paragraph(caption, style="Caption")
    table = doc.add_table(rows=1, cols=len(headings))
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, headings):
        cell.text = text
        cell.paragraphs[0].runs[0].bold = True
    for row in rows + ([total] if total else []):
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = str(value)
    if total:
        for cell in table.rows[-1].cells:
            for run in cell.paragraphs[0].runs:
                run.bold = True
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(10)
    doc.add_paragraph()


def template(production: Production, kind: str) -> bytes:
    """A Word template for one uploaded part, prefilled with what the system knows."""
    conference = production.conference
    number = conference.number
    editors = editors_of(production)
    if kind == BookPart.Kind.FOREWORD:
        stats = statistics(production)
        doc = _document("Foreword")
        doc.add_paragraph(
            f"[Opening: the {ordinal(number)} Annual Conference of the International Group for Lean Construction "
            f"(IGLC{number}) in {place_of(conference) or '…'}, {date_line(conference)}.]")
        doc.add_paragraph(
            f"This year’s proceedings contain {stats['papers']} papers from authors affiliated with institutions in "
            f"{len(stats['countries'])} countries, as detailed in Table 1. [Themes of this year’s papers.]")
        _table(doc, "Table 1 Papers published per country", ["Country of the first author’s institution",
                                                             "Published papers"],
               [[c or "Unknown", n] for c, n in stats["countries"]], ["TOTAL", stats["papers"]])
        doc.add_paragraph("[The tracks and the review process: papers submitted and accepted per track are in "
                          "Table 2; fill in the submitted papers from ConfTool.]")
        _table(doc, "Table 2 Papers per track", ["Track", "Papers submitted", "Papers accepted"],
               [[t, "", n] for t, n in stats["tracks"]], ["TOTAL", "", stats["papers"]])
        chairs = production.track_chairs.select_related("track")
        _table(doc, "Table 3 Track chairs", ["Track", "Track chair name and affiliation"],
               [[c.track.title, f"{c.name}\n{c.affiliation}".strip()] for c in chairs] or [["", ""]])
        doc.add_paragraph("[Thanks to the track chairs, the reviewers, the local organisers and the conference "
                          "chair" + (f", {production.conference_chair}" if production.conference_chair else "") + ".]")
        signature = doc.add_paragraph()
        signature.add_run(names_joined(editors) or "[Editors]").bold = True
        doc.add_paragraph(f"Editors and Scientific Chairs of IGLC{number}").runs[0].italic = True
    elif kind == BookPart.Kind.ORGANISATION:
        doc = _document("Conference organisation")
        doc.add_heading("Editorial team", level=2)
        for name in editors or ["[Name], Editor"]:
            doc.add_paragraph(f"{name}, Editor")
        doc.add_heading("Local organising committee", level=2)
        doc.add_paragraph(f"{production.conference_chair or '[Name]'}, Chair")
        doc.add_paragraph("[Name]")
        doc.add_heading("Conference organisers", level=2)
        doc.add_paragraph("[Host institution(s)]")
        doc.add_paragraph("International Group for Lean Construction")
    elif kind == BookPart.Kind.REVIEWERS:
        doc = _document("List of reviewers")
        doc.add_paragraph("[One row per reviewer, sorted by last name: “Last name, First name” and the "
                          "affiliation with country. ConfTool’s list of reviewers can be pasted in.]")
        _table(doc, "", ["Reviewer", "Affiliation"], [["Surname, Given name", "University, Country"]])
    elif kind == BookPart.Kind.MESSAGE:
        doc = _document("Message from the conference chair")
        doc.add_paragraph("[Text. The heading can be changed, e.g. “Message from the host institution”.]")
        signature = doc.add_paragraph()
        signature.add_run(production.conference_chair or "[Name]").bold = True
        doc.add_paragraph(f"Conference Chair, IGLC{number}").runs[0].italic = True
    elif kind == BookPart.Kind.SPONSORS:
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = _document("Sponsors")
        doc.add_paragraph("[Thanks to the sponsors (optional).]")
        for level in ("Platinum sponsors", "Gold sponsors", "Silver sponsors", "Bronze sponsors",
                      "Supporting partners"):
            doc.add_heading(level, level=2)
            logos = doc.add_paragraph("[Logos: Insert → Pictures, “In line with text”, at most 4 cm high. "
                                      "Delete the levels not used.]")
            logos.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        doc = _document("[Heading]")
        doc.add_paragraph("[Text. Use the styles of this template only: Heading 1 for the heading, Heading 2 "
                          "for subheadings, Normal for text, Caption for table and figure captions.]")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------- the papers' PDFs

def _cache_name(submission: Submission) -> str:
    return f"production/iglc{submission.production.conference.number}/book/papers/{submission.conftool_id:04d}.pdf"


def paper_pdf(submission: Submission) -> bytes | None:
    """The published PDF of a paper, if it is at hand (made here, or fetched before)."""
    storage = private_storage()
    if submission.published_pdf:
        with submission.published_pdf.open("rb") as handle:
            return handle.read()
    name = _cache_name(submission)
    if storage.exists(name):
        with storage.open(name, "rb") as handle:
            return handle.read()
    return None


def missing_pdfs(production: Production) -> list[Submission]:
    storage = private_storage()
    return [s for _, subs in placed_papers(production) for s in subs
            if not s.published_pdf and not storage.exists(_cache_name(s))]


def fetch_next(production: Production, n: int = 10, after: int = 0) -> dict:
    """Download up to n published PDFs (papers published before the production tools) into the
    private files, checking that each has as many pages as its page range. Goes through the
    papers once, in ConfTool ID order after `after`, so a paper that fails is not retried in
    the same round; returns the last ID tried as `next`."""
    from pypdf import PdfReader

    todo = sorted((s for s in missing_pdfs(production) if s.conftool_id > after), key=lambda s: s.conftool_id)
    batch, problems, fetched = todo[:n], [], 0
    for submission in batch:
        paper = submission.paper
        if not paper.full_text_url:
            problems.append(f"{submission.conftool_id}: no PDF in the archive")
            continue
        try:
            request = urllib.request.Request(paper.full_text_url, headers={"User-Agent": "iglc.net proceedings"})
            with urllib.request.urlopen(request, timeout=60) as response:
                data = response.read()
            pages = len(PdfReader(io.BytesIO(data)).pages)
        except Exception as error:  # noqa: BLE001 - reported to the editor
            problems.append(f"{submission.conftool_id}: could not be downloaded or read ({error})")
            continue
        expected = paper.last_page - paper.first_page + 1
        if pages != expected:
            problems.append(f"{submission.conftool_id}: the PDF has {pages} pages, the page range {paper.pages} "
                            f"has {expected}")
        private_storage().save(_cache_name(submission), ContentFile(data))
        fetched += 1
    return {"fetched": fetched, "left": len(todo) - len(batch), "problems": problems,
            "next": batch[-1].conftool_id if batch else after}


# ---------------------------------------------------------------- generated pages

def _fonts():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    from .publish import fonts_folder

    folder = fonts_folder()
    for name, file in (("IGLC-Regular", "times.ttf"), ("IGLC-Italic", "timesi.ttf"), ("IGLC-Bold", "timesbd.ttf")):
        try:
            pdfmetrics.getFont(name)
        except KeyError:
            if not (folder / file).exists():
                raise BookError(f"The font {file} is missing (with times.ttf and timesi.ttf)")
            pdfmetrics.registerFont(TTFont(name, str(folder / file)))


@dataclass
class Link:
    page: int            # page of the generated part (0-based)
    rect: tuple
    target: Submission


class _Pages:
    """A small layout helper on a reportlab canvas: text lines top-down, new pages as needed."""

    def __init__(self):
        from reportlab.pdfgen import canvas

        self.buffer = io.BytesIO()
        self.canvas = canvas.Canvas(self.buffer, pagesize=PAGE)
        self.page, self.y, self.links = 0, PAGE[1] - MARGIN, []
        self.width = PAGE[0] - 2 * MARGIN

    def need(self, height):
        if self.y - height < MARGIN:
            self.new_page()

    def new_page(self):
        self.canvas.showPage()
        self.page += 1
        self.y = PAGE[1] - MARGIN

    def text(self, text, font="IGLC-Regular", size=11, leading=None, align="left", width=None, x=None):
        from reportlab.lib.utils import simpleSplit

        leading = leading or size * 1.25
        width = width or self.width
        x = MARGIN if x is None else x
        for line in simpleSplit(text, font, size, width) or [""]:
            self.need(leading)
            self.y -= leading
            self.canvas.setFont(font, size)
            if align == "center":
                self.canvas.drawCentredString(PAGE[0] / 2, self.y, line)
            else:
                self.canvas.drawString(x, self.y, line)

    def space(self, height):
        self.y -= height

    def finish(self) -> bytes:
        self.canvas.showPage()
        self.canvas.save()
        return self.buffer.getvalue()


def colophon(production: Production) -> bytes:
    conference, editors = production.conference, editors_of(production)
    year = conference.year
    title = conference.proceedings_title or f"Proceedings of the {ordinal(conference.number)} Annual Conference " \
                                             f"of the International Group for Lean Construction (IGLC {conference.number})"
    holders = production.copyright_holders or names_joined(editors)
    p = _Pages()
    p.text(title, "IGLC-Bold", 10)
    p.space(6)
    p.text(f"{names_joined(editors)} (editors)", size=10)
    p.space(12)
    p.text(f"© {year} {holders}", size=10)
    p.space(24)
    p.text(f"The IGLC{conference.number} Proceedings Editors received copyright permission from the authors to "
           "publish these Proceedings of the "
           f"{ordinal(conference.number)} Annual Conference of the International Group for Lean Construction in hard "
           "copy format as well as in digital format for online posting without access restrictions. Other than this "
           "copyright transferred to the Proceedings Editors, the authors reserve all proprietary rights (such as "
           "patent rights) in their work. The authors have the right to republish their work, in whole or part, in any "
           "publication of which they are an author or editor, and to make other personal use of the work. Any such "
           "republication or personal use must explicitly identify prior publication, including the names of the "
           "Proceedings Editors, and the page numbers in these Proceedings.", size=9, leading=11.5)
    p.y = PAGE[1] / 2 + 40
    p.text(f"Published {year} by the International Group for Lean Construction", size=10)
    p.text("www.iglc.net", size=10)
    p.space(24)
    for label, value in (("ISSN", production.issn_print and f"{production.issn_print} (printed)"),
                         ("ISSN", production.issn_electronic and f"{production.issn_electronic} (electronic)"),
                         ("ISBN", production.isbn_print and f"{production.isbn_print} (printed)"),
                         ("ISBN", production.isbn_pdf and f"{production.isbn_pdf} (PDF)")):
        if value:
            p.text(f"{label}: {value}", size=10)
    return p.finish()


def title_page(production: Production) -> bytes:
    conference, editors = production.conference, editors_of(production)
    p = _Pages()
    p.text(f"Proceedings of IGLC{conference.number}", size=14, align="center")
    p.space(70)
    for line in (f"{ordinal(conference.number)} ANNUAL CONFERENCE", "of the", "INTERNATIONAL GROUP", "for",
                 "LEAN CONSTRUCTION"):
        p.text(line, size=16, leading=28, align="center")
    p.space(60)
    p.text("Edited by", size=12, leading=18, align="center")
    for name in editors:
        p.text(name, size=12, leading=18, align="center")
    if production.conference_chair:
        p.space(30)
        p.text("Conference Chair", size=12, leading=18, align="center")
        for name in production.conference_chair.split(";"):
            p.text(name.strip(), size=12, leading=18, align="center")
    p.space(40)
    place = [conference.city] if conference.city == conference.country else [conference.city, conference.country]
    for line in (*place, date_line(conference)):
        if line:
            p.text(line, size=12, leading=18, align="center")
    return p.finish()


def _leader_line(p: _Pages, lines: list[str], number: str, font="IGLC-Bold", size=10):
    """Title lines; the last one with dot leaders to the page number at the right margin."""
    from reportlab.pdfbase.pdfmetrics import stringWidth

    for index, line in enumerate(lines):
        p.need(13)
        p.y -= 13
        p.canvas.setFont(font, size)
        p.canvas.drawString(MARGIN, p.y, line)
        if index == len(lines) - 1:
            right = PAGE[0] - MARGIN
            p.canvas.drawRightString(right, p.y, number)
            start = MARGIN + stringWidth(line + " ", font, size)
            end = right - stringWidth(" " + number, font, size)
            dot = stringWidth(".", font, size)
            if end > start:
                p.canvas.drawString(start, p.y, "." * int((end - start) / dot))


def contents(production: Production, first_page_of: dict) -> tuple[bytes, list[Link]]:
    """The table of contents; links point at papers (resolved to pages when joined)."""
    from reportlab.lib.utils import simpleSplit

    chairs = {}
    for chair in production.track_chairs.all():
        chairs.setdefault(chair.track_id, []).append(chair.name)
    p = _Pages()
    p.text("TABLE OF CONTENTS", "IGLC-Bold", 12, leading=16)
    p.space(10)
    for track, subs in placed_papers(production):
        p.need(60)
        p.space(10)
        p.text((track.title if track else "Other papers").upper(), "IGLC-Bold", 10, leading=13)
        if track and chairs.get(track.pk):
            p.text("Track chair" + ("s" if len(chairs[track.pk]) > 1 else "") + ": " + ", ".join(chairs[track.pk]),
                   "IGLC-Italic", 9, leading=11)
        p.space(6)
        for submission in subs:
            paper = submission.paper
            if not paper:
                continue
            lines = simpleSplit(paper.title, "IGLC-Bold", 10, p.width - 40)
            authors = simpleSplit(", ".join(f"{a.first_name} {a.last_name}".strip() for a in paper.authors.all()),
                                  "IGLC-Italic", 9, p.width - 40)
            p.need(13 * len(lines) + 11 * len(authors) + 8)
            top, page = p.y, p.page
            _leader_line(p, lines, str(paper.first_page))
            for line in authors:
                p.y -= 11
                p.canvas.setFont("IGLC-Italic", 9)
                p.canvas.drawString(MARGIN, p.y, line)
            p.links.append(Link(page, (MARGIN, p.y - 2, PAGE[0] - MARGIN, top), submission))
            p.space(8)
    return p.finish(), p.links


def _sort_key(text: str) -> str:
    return unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode().lower()


def author_index(production: Production) -> tuple[bytes, list[Link]]:
    """Every author, with the first page of each of their papers, in two columns."""
    from reportlab.lib.utils import simpleSplit
    from reportlab.pdfbase.pdfmetrics import stringWidth

    entries = {}
    for _, subs in placed_papers(production):
        for submission in subs:
            if not submission.paper:
                continue
            for author in submission.paper.authors.all():
                name = f"{author.last_name}, {author.first_name}".strip(", ")
                entries.setdefault((_sort_key(name), name), []).append(submission)
    p = _Pages()
    p.text("AUTHOR INDEX", "IGLC-Bold", 12, leading=16)
    p.space(10)
    column_width = (p.width - 20) / 2
    top, column = p.y, 0
    for (_, name), subs in sorted(entries.items()):
        pages = sorted({s.paper.first_page for s in subs})
        numbers = ", ".join(map(str, pages))
        lines = simpleSplit(f"{name}  {numbers}", "IGLC-Regular", 9, column_width - 10)
        height = 11 * len(lines)
        if p.y - height < MARGIN:
            if column == 0:
                column, p.y = 1, top
            else:
                p.new_page()
                top, column = p.y, 0
        x = MARGIN + column * (column_width + 20)
        for index, line in enumerate(lines):
            p.y -= 11
            p.canvas.setFont("IGLC-Regular", 9)
            p.canvas.drawString(x + (10 if index else 0), p.y, line)
        # each page number links to its paper
        width = stringWidth(lines[-1], "IGLC-Regular", 9)
        numbers_x = x + (10 if len(lines) > 1 else 0) + width - stringWidth(numbers, "IGLC-Regular", 9)
        for page_number, submission in zip(pages, sorted(subs, key=lambda s: s.paper.first_page)):
            w = stringWidth(str(page_number), "IGLC-Regular", 9)
            if numbers_x >= x:
                p.links.append(Link(p.page, (numbers_x, p.y - 2, numbers_x + w, p.y + 8), submission))
            numbers_x += w + stringWidth(", ", "IGLC-Regular", 9)
    return p.finish(), p.links


def _numbers_overlay(pages: list[tuple[tuple, str]]) -> bytes:
    """One page per entry with the page number centred at the bottom ('' for none)."""
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    c = canvas.Canvas(buffer)
    for size, label in pages:
        c.setPageSize(size)
        if label:
            c.setFont("IGLC-Regular", 10)
            c.drawCentredString(size[0] / 2, 36, label)
        c.showPage()
    c.save()
    return buffer.getvalue()


# ---------------------------------------------------------------- checking uploaded parts

ALLOWED_FONTS = re.compile(r"times|symbol|wingding", re.I)


def _page_fonts(page) -> dict[str, str]:
    """Resource name -> base font of the fonts the page's own text uses (not those in figures)."""
    fonts = {}
    resources = page.get("/Resources") or {}
    for name, ref in (resources.get("/Font") or {}).items():
        font = ref.get_object()
        fonts[str(name)] = str(font.get("/BaseFont", "")).lstrip("/").split("+")[-1]
    return fonts


def check_part(data: bytes, kind: str) -> list[str]:
    """Problems with an uploaded part: covers are one A4 page; sections must follow the IGLC
    template: A4, Times New Roman, text within the margins, header and footer empty (the page
    numbers go there). Logos and figures are pictures and not checked."""
    from pypdf import PdfReader

    from .pdf_running import FOOTER_ZONE, HEADER_ZONE, _text_positions

    try:
        reader = PdfReader(io.BytesIO(data))
        pages = list(reader.pages)
    except Exception:  # noqa: BLE001
        return ["This is not a PDF that can be read."]
    problems = []
    if any(abs(float(p.mediabox.width) - PAGE[0]) > 3 or abs(float(p.mediabox.height) - PAGE[1]) > 3 for p in pages):
        problems.append("The pages must be A4 (21 × 29.7 cm).")
    if kind in BookPart.COVERS:
        if len(pages) != 1:
            problems.append("A cover is one page.")
        return problems
    fonts, outside, running = Counter(), set(), set()
    for number, page in enumerate(pages, 1):
        names = _page_fonts(page)
        width, height = float(page.mediabox.width), float(page.mediabox.height)
        for x, y, text, font in _text_positions(page, reader, with_x=True):
            base = names.get(font, font)
            if base and not ALLOWED_FONTS.search(base):
                fonts[base] += len(text)
            if y > height - HEADER_ZONE or y < FOOTER_ZONE:
                running.add(number)
            elif x < MARGIN - 8 or x > width - MARGIN:
                outside.add(number)
    fonts = sorted(f for f, characters in fonts.items() if characters >= 20)  # a stray symbol is tolerated
    if fonts:
        problems.append("Only Times New Roman may be used (the template's styles); found: " + ", ".join(fonts) + ".")
    if running:
        problems.append("The header and footer must be empty – the page numbers are printed there (page "
                        + ", ".join(map(str, sorted(running))) + ").")
    if outside:
        problems.append("Text outside the template's margins (page " + ", ".join(map(str, sorted(outside))) + ").")
    return problems


# ---------------------------------------------------------------- joining

def parts_of(production: Production) -> dict[str, list[BookPart]]:
    parts = {}
    for part in production.book_parts.all():
        parts.setdefault(part.kind, []).append(part)
    return parts


def problems(production: Production, final: bool = False) -> list[str]:
    """What stops the book from being made (final: what stops it from being published)."""
    found = []
    missing = missing_pdfs(production)
    if missing:
        found.append(f"{len(missing)} papers' PDFs are not fetched yet.")
    if not editors_of(production):
        found.append("The conference has no editors in the archive (Manage → Conferences).")
    try:
        _fonts()
    except Exception as error:  # noqa: BLE001
        found.append(str(error))
    if final:
        parts = parts_of(production)
        for kind in (BookPart.Kind.COVER, BookPart.Kind.FOREWORD):
            if kind not in parts:
                found.append(f"No {BookPart.Kind(kind).label.lower()} uploaded.")
        if not (production.isbn_pdf or production.isbn_print):
            found.append("No ISBN: the publisher enters it.")
    return found


def _read(part_or_bytes):
    from pypdf import PdfReader

    if isinstance(part_or_bytes, bytes):
        return PdfReader(io.BytesIO(part_or_bytes))
    with part_or_bytes.pdf.open("rb") as handle:
        return PdfReader(io.BytesIO(handle.read()))


def build(production: Production) -> bytes:
    """Join everything into the full proceedings PDF."""
    from pypdf import PdfReader, PdfWriter
    from pypdf.annotations import Link as LinkAnnotation
    from pypdf.generic import Fit

    found = problems(production)
    if found:
        raise BookError(found[0])
    _fonts()
    parts = parts_of(production)
    writer = PdfWriter()
    numbers = []  # (size, label) per page, for the page-number overlay
    outline = []  # (title, page index, parent title or None)

    def add(reader, label_of, title=None, parent=None):
        start = len(writer.pages)
        for index, page in enumerate(reader.pages):
            writer.add_page(page)
            size = (float(page.mediabox.width), float(page.mediabox.height))
            numbers.append((size, label_of(index)))
        if title:
            outline.append((title, start, parent))
        return start

    def blank(label=""):
        writer.add_blank_page(*PAGE)
        numbers.append((PAGE, label))

    # cover and colophon: no numbers
    if parts.get(BookPart.Kind.COVER):
        add(_read(parts[BookPart.Kind.COVER][0]), lambda i: "", "Cover")
    else:
        blank()
    add(PdfReader(io.BytesIO(colophon(production))), lambda i: "")
    if len(writer.pages) % 2:
        blank()
    front_start = len(writer.pages)

    base = lambda: len(writer.pages) - front_start + 1  # noqa: E731
    b = base()
    add(PdfReader(io.BytesIO(title_page(production))), lambda i, b=b: roman(b + i), "Title page")
    sections = [p for p in production.book_parts.all() if p.kind not in BookPart.COVERS]
    for part in [p for p in sections if p.placement == BookPart.Placement.FRONT]:
        b = base()
        add(_read(part), lambda i, b=b: roman(b + i), part.label)
    toc_pdf, toc_links = contents(production, {})
    b = base()
    toc_start = add(PdfReader(io.BytesIO(toc_pdf)), lambda i, b=b: roman(b + i), "Table of contents")
    if len(writer.pages) % 2:
        blank()  # page 1 is a right-hand page
    # the papers, as published (their own page numbers)
    paper_start = len(writer.pages)
    starts = {}
    for track, subs in placed_papers(production):
        track_title = track.title if track else "Other papers"
        first_of_track = True
        for submission in subs:
            data = paper_pdf(submission)
            if data is None:
                raise BookError(f"Paper {submission.conftool_id}'s PDF is missing")
            starts[submission.pk] = add(PdfReader(io.BytesIO(data)), lambda i: "", submission.paper.title,
                                        track_title)
            if first_of_track:
                outline.append((track_title, starts[submission.pk], None))
                first_of_track = False
    last_paper_page = max(s.paper.last_page for _, subs in placed_papers(production) for s in subs if s.paper)
    # author index, continuing the numbering
    index_first_number = last_paper_page + 1
    if len(writer.pages) % 2:
        blank()  # the index starts on a right-hand page; the blank page is not numbered
        index_first_number += 1
    index_pdf, index_links = author_index(production)
    index_start = add(PdfReader(io.BytesIO(index_pdf)), lambda i: str(index_first_number + i), "Author index")
    next_number = index_first_number + (len(writer.pages) - index_start)
    for part in [p for p in sections if p.placement == BookPart.Placement.BACK]:
        if len(writer.pages) % 2:
            blank()  # each back section starts on a right-hand page (the blank one is not numbered)
            next_number += 1
        start = add(_read(part), lambda i, f=next_number: str(f + i), part.label)
        next_number += len(writer.pages) - start
    if parts.get(BookPart.Kind.BACK_COVER):
        if len(writer.pages) % 2 == 0:
            blank()  # the back cover is a left-hand page
        add(_read(parts[BookPart.Kind.BACK_COVER][0]), lambda i: "", "Back cover")
    # page numbers on the front matter and the index
    overlay = PdfReader(io.BytesIO(_numbers_overlay(numbers)))
    for page, layer, (_, label) in zip(writer.pages, overlay.pages, numbers):
        if label:
            page.merge_page(layer)
    # links from the contents and the index to the papers
    for start, links in ((toc_start, toc_links), (index_start, index_links)):
        for link in links:
            if link.target.pk in starts:
                writer.add_annotation(start + link.page, LinkAnnotation(rect=link.rect, target_page_index=starts[
                    link.target.pk]))
    # bookmarks
    parents = {}
    for title, page, parent in outline:
        if parent is None:
            parents[title] = writer.add_outline_item(title, page, fit=Fit.fit())
    for title, page, parent in outline:
        if parent is not None:
            writer.add_outline_item(title, page, parent=parents.get(parent), fit=Fit.fit())
    # page labels, so a viewer shows the printed numbers
    writer.set_page_label(0, front_start - 1, prefix="Cover ")
    writer.set_page_label(front_start, paper_start - 1, style="/r")
    writer.set_page_label(paper_start, len(writer.pages) - 1, style="/D", start=production.first_page)
    conference = production.conference
    writer.add_metadata({"/Title": conference.proceedings_title or f"Proceedings IGLC{conference.number}",
                         "/Author": names_joined(editors_of(production)),
                         "/Subject": f"Proceedings of the {ordinal(conference.number)} Annual Conference of the "
                                     "International Group for Lean Construction"})
    writer.compress_identical_objects(remove_identicals=True, remove_orphans=True)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def draft_name(production: Production) -> str:
    return f"production/iglc{production.conference.number}/book/draft.pdf"


def save_draft(production: Production) -> int:
    data = build(production)
    storage = private_storage()
    if storage.exists(draft_name(production)):
        storage.delete(draft_name(production))
    storage.save(draft_name(production), ContentFile(data))
    return len(data)


def publish_book(production: Production, user=None) -> str:
    """Build the final book (with the ISBNs), make it public, and link it from the conference page."""
    from apps.archive.models import ProceedingsFile, Volume

    found = problems(production, final=True)
    if found:
        raise BookError(found[0])
    data = build(production)
    conference = production.conference
    name = default_storage.save(f"proceedings/IGLC{conference.number}-Proceedings.pdf", ContentFile(data))
    if production.book and production.book.name != name:
        default_storage.delete(production.book.name)
    production.book.name = name
    production.status = Production.Status.COMPLETE
    production.save(update_fields=["book", "status"])
    from .publish import public_url

    url = public_url(name)
    ProceedingsFile.objects.update_or_create(conference=conference, label="Full proceedings",
                                             defaults={"url": url, "order": 1})
    papers = [s.paper for _, subs in placed_papers(production) for s in subs if s.paper]
    volume, _ = Volume.objects.update_or_create(
        conference=conference, number=1,
        defaults={"first_page": min(p.first_page for p in papers), "last_page": max(p.last_page for p in papers),
                  "isbn": production.isbn_pdf or production.isbn_print})
    conference.papers.filter(pk__in=[p.pk for p in papers]).update(volume=volume)
    return url
