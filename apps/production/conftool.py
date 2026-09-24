"""Reading ConfTool's export of accepted papers (Excel or CSV).

ConfTool's export columns vary with the conference's settings, so the columns are found by
their headings (case and punctuation ignored), and can be named explicitly. One row per paper.
Several authors in one cell are separated by semicolons or line breaks, or numbered
("Author 1 Name", "Author 2 Name", ...).
"""

from __future__ import annotations

import csv
import io
import re
from pathlib import Path

COLUMNS = {
    "id": ["paper id", "paperid", "id", "submission id", "contribution id", "conftool id"],
    "title": ["title", "paper title", "contribution title"],
    "track": ["track", "topic", "topics", "session track", "contribution type track", "subject area"],
    "status": ["acceptance status", "status", "decision"],
    "authors": ["authors", "author names", "all authors"],
    "organisations": ["organisations", "organizations", "affiliations", "institutions"],
    "emails": ["emails", "email addresses", "author emails", "e mails"],
}


def _norm(heading) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(heading or "").lower()).strip()


def read_table(path) -> list[dict]:
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        sheet = load_workbook(path, read_only=True, data_only=True).active
        rows = list(sheet.iter_rows(values_only=True))
        headings = [str(h or "") for h in rows[0]]
        return [dict(zip(headings, ["" if v is None else str(v) for v in row])) for row in rows[1:] if any(row)]
    text = path.read_bytes().decode("utf-8-sig", errors="replace")
    dialect = csv.Sniffer().sniff(text[:5000], delimiters=",;\t")
    return list(csv.DictReader(io.StringIO(text), dialect=dialect))


def find_columns(headings, explicit: dict | None = None) -> dict:
    by_norm = {_norm(h): h for h in headings}
    found = {}
    for key, names in COLUMNS.items():
        if explicit and explicit.get(key):
            found[key] = explicit[key]
            continue
        found[key] = next((by_norm[n] for n in names if n in by_norm), None)
    return found


def _split(value) -> list[str]:
    return [v.strip() for v in re.split(r"\s*(?:;|\n)\s*", value or "") if v.strip()]


def authors_of(row: dict, columns: dict) -> list[dict]:
    numbered = sorted({int(m.group(1)) for h in row for m in [re.match(r"author\s*(\d+)", _norm(h))] if m})
    if numbered:
        authors = []
        for n in numbered:
            def cell(*words, n=n):
                for heading, value in row.items():
                    h = _norm(heading)
                    if re.match(rf"author\s*{n}\b", h) and all(w in h for w in words):
                        return str(value or "").strip()
                return ""
            name = cell("name") or " ".join(filter(None, [cell("first"), cell("last")]))
            if name:
                authors.append({"name": name, "organisation": cell("organi") or cell("affiliation"),
                                "email": cell("mail")})
        return authors
    names = _split(row.get(columns.get("authors") or "", ""))
    if len(names) == 1 and "," in names[0]:
        names = [n.strip() for n in re.split(r",\s*|\s+and\s+|\s*&\s*", names[0]) if n.strip()]
    organisations = _split(row.get(columns.get("organisations") or "", ""))
    emails = _split(row.get(columns.get("emails") or "", ""))
    return [{"name": name, "organisation": organisations[i] if i < len(organisations) else "",
             "email": emails[i] if i < len(emails) else ""} for i, name in enumerate(names)]


def read_accepted(path, explicit: dict | None = None, accepted_only=True):
    rows = read_table(path)
    if not rows:
        return [], {}
    columns = find_columns(rows[0].keys(), explicit)
    if not columns["id"] or not columns["title"]:
        raise ValueError(f"No ID or title column found among: {', '.join(rows[0].keys())}")
    papers = []
    for row in rows:
        status = _norm(row.get(columns["status"] or "", ""))
        if accepted_only and columns["status"] and status and "accept" not in status:
            continue
        digits = re.sub(r"\D", "", str(row.get(columns["id"], "")))
        if not digits:
            continue
        papers.append({
            "conftool_id": int(digits),
            "title": str(row.get(columns["title"], "")).strip(),
            "track": str(row.get(columns["track"] or "", "")).strip(),
            "authors": authors_of(row, columns),
        })
    return papers, columns
