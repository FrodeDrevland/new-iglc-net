"""The programme as a spreadsheet (Excel): exported in the same form that can be imported, so a
programme drafted in Excel can be brought in and then adjusted in the builder.

One row per session. Columns (found by their headings, in any order; only Day, Start and End are
needed): Day, Start, End, Lane, Room, Part, Kind, Code, Title, Plenary, Chairs, Papers, Contributions, Notes.

- Day: a date (2027-07-20) or a cell formatted as a date. Start and End: 10:30.
- Lane: which parallel lane (1, 2, 3 ...) the session is drawn in; blank: the first free one.
- Room: its name, or blank to give it later; a new name adds the room (for those who keep the rooms).
- Part: only for a day that holds two parts; otherwise the day's own. Kind: as in the builder,
  e.g. "Paper session".
- Plenary: yes or no. Chairs: "Ann Smith (NTNU); Bo Jones".
- Papers: ConfTool IDs, e.g. "101, 117, 123". Contributions: titles, separated by ";". A title not in
  the list of contributions is added to it.

A row whose day, start and lane, or day, start and room, or day and code match a session updates it;
other rows add sessions.
"""

from __future__ import annotations

import io
import re
from datetime import date as Date, datetime, time

from django.db import transaction

from . import access
from .builder import BuildError, _save, place
from .models import Contribution, Location, Session, SessionPerson

HEADINGS = ["Day", "Start", "End", "Lane", "Room", "Part", "Kind", "Code", "Title", "Plenary", "Chairs", "Papers",
            "Contributions", "Notes"]
YES = {"yes", "y", "true", "1", "x", "ja"}


def export(programme) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    book = Workbook()
    sheet = book.active
    sheet.title = f"IGLC {programme.conference.number} programme"
    sheet.append(HEADINGS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sessions = (programme.sessions.select_related("location", "part")
                .prefetch_related("people", "items__submission", "items__contribution"))
    for s in sessions:
        chairs = "; ".join(f"{p.name} ({p.affiliation})" if p.affiliation else p.name
                           for p in s.people.all() if p.role in ("chair", "co_chair"))
        papers = ", ".join(str(i.submission.conftool_id) for i in s.items.all() if i.submission_id)
        contributions = "; ".join(i.contribution.title for i in s.items.all() if i.contribution_id)
        sheet.append([s.date, s.start.strftime("%H:%M"), s.end.strftime("%H:%M"), s.lane or "",
                      s.location.name if s.location_id else "", s.part.name, s.get_kind_display(), s.code,
                      s.title, "yes" if s.plenary else "", chairs, papers, contributions, s.notes])
    for row in sheet.iter_rows(min_row=2, max_col=1):
        row[0].number_format = "yyyy-mm-dd"
    for column, width in zip("ABCDEFGHIJKLMN", (12, 7, 7, 6, 18, 20, 16, 7, 40, 8, 30, 20, 30, 30)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _norm(text) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(text or "").lower()).strip()


def _day(value) -> Date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, Date):
        return value
    text = str(value or "").strip()
    for pattern in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    raise BuildError(f"Not a day: “{text}” (write it as 2027-07-20)")


def _clock(value) -> time:
    if isinstance(value, datetime):
        return value.time().replace(second=0, microsecond=0)
    if isinstance(value, time):
        return value.replace(second=0, microsecond=0)
    if isinstance(value, float) and 0 <= value < 1:  # Excel's fraction of a day
        minutes = round(value * 24 * 60)
        return time(minutes // 60, minutes % 60)
    match = re.match(r"^\s*(\d{1,2})[:.](\d{2})", str(value or ""))
    if not match:
        raise BuildError(f"Not a time: “{value}” (write it as 10:30)")
    return time(int(match.group(1)), int(match.group(2)))


def _rows(path):
    from openpyxl import load_workbook

    sheet = load_workbook(path, read_only=True, data_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    columns = {}
    for index, heading in enumerate(rows[0]):
        for name in HEADINGS:
            if _norm(heading) == _norm(name):
                columns[name] = index
    missing = [name for name in ("Day", "Start", "End") if name not in columns]
    if missing:
        raise BuildError("The sheet needs the columns " + ", ".join(missing) + ".")
    return [(number, {name: row[i] if i < len(row) else None for name, i in columns.items()})
            for number, row in enumerate(rows[1:], start=2) if any(v not in (None, "") for v in row)]


@transaction.atomic
def import_sheet(programme, user, path, replace_days: bool = False) -> dict:
    parts = access.editable_parts(user, programme)
    if not parts:
        raise BuildError("You do not edit any part of the programme.")
    by_part = {}
    for part in parts:
        by_part[_norm(part.name)] = part
        by_part.setdefault(_norm(part.get_kind_display()), part)
        by_part.setdefault(_norm(part.kind), part)
    kinds = {_norm(label): key for key, label in Session.Kind.choices} | {key: key for key in Session.Kind.values}
    kinds.update({"paper session": "papers", "papers": "papers", "posters": "posters", "lunch": "meal",
                  "dinner": "social", "coffee": "break", "coffee break": "break"})
    rooms = {_norm(loc.name): loc for loc in programme.locations.all()}
    can_rooms = access.can_edit_locations(user, programme)
    rows = _rows(path)
    added = updated = 0
    problems = []
    if replace_days:
        days = set()
        for number, row in rows:
            try:
                days.add(_day(row.get("Day")))
            except BuildError:
                pass
        programme.sessions.filter(date__in=days, part__in=parts).delete()
    for number, row in rows:
        try:
            with transaction.atomic():
                day, start, end = _day(row.get("Day")), _clock(row.get("Start")), _clock(row.get("End"))
                room = None
                room_name = str(row.get("Room") or "").strip()
                if room_name:
                    room = rooms.get(_norm(room_name))
                    if room is None:
                        if not can_rooms:
                            raise BuildError(f"No room “{room_name}” (the organisers add rooms)")
                        room = Location.objects.create(programme=programme, name=room_name[:200],
                                                       sort_order=len(rooms) + 1)
                        rooms[_norm(room_name)] = room
                part_name = str(row.get("Part") or "").strip()
                if part_name:
                    part = by_part.get(_norm(part_name))
                    if part is None:
                        raise BuildError(f"No part “{part_name}” that you edit")
                else:  # the day's own part
                    from .builder import _day_part

                    part = _day_part(user, programme, day)
                kind = kinds.get(_norm(row.get("Kind")), "papers" if not row.get("Kind") else None)
                if kind is None:
                    raise BuildError(f"Unknown kind “{row.get('Kind')}”")
                code = str(row.get("Code") or "").strip()
                lane_text = str(row.get("Lane") or "").strip()
                lane = int(float(lane_text)) if lane_text.replace(".", "", 1).isdigit() else None
                existing = None
                if lane:
                    existing = programme.sessions.filter(date=day, start=start, lane=lane, part__in=parts).first()
                if existing is None and room is not None:
                    existing = programme.sessions.filter(date=day, start=start, location=room, part__in=parts).first()
                if existing is None and code:
                    existing = programme.sessions.filter(date=day, code=code, part__in=parts).first()
                session = existing or Session(programme=programme)
                session.part, session.date, session.start, session.end = part, day, start, end
                session.location, session.kind, session.code = room, kind, code
                if lane:
                    session.lane = lane
                session.title = str(row.get("Title") or "").strip()[:300]
                session.plenary = _norm(row.get("Plenary")) in YES
                if row.get("Notes") is not None:
                    session.notes = str(row.get("Notes") or "").strip()
                _save(session)
                if row.get("Chairs"):
                    session.people.filter(role__in=["chair", "co_chair"]).delete()
                    for order, chunk in enumerate(re.split(r"\s*;\s*", str(row["Chairs"])), start=1):
                        match = re.match(r"^(.*?)\s*(?:\((.*)\))?$", chunk.strip())
                        if match and match.group(1):
                            SessionPerson.objects.create(session=session, name=match.group(1)[:200],
                                                         affiliation=(match.group(2) or "")[:300], order=order)
                entries = []
                for number_text in re.findall(r"\d+", str(row.get("Papers") or "")):
                    submission = programme.conference.production.submissions.filter(conftool_id=int(number_text)).first() \
                        if hasattr(programme.conference, "production") else None
                    if submission is None:
                        raise BuildError(f"No paper {number_text}")
                    entries.append(f"s{submission.pk}")
                for title in [t.strip() for t in str(row.get("Contributions") or "").split(";") if t.strip()]:
                    contribution = programme.contributions.filter(title__iexact=title).first() or \
                        Contribution.objects.create(programme=programme, part=part, title=title[:300])
                    entries.append(f"c{contribution.pk}")
                for index, entry in enumerate(entries):
                    place(programme, user, {"session": session.pk, "entry": entry, "index": index})
                if existing:
                    updated += 1
                else:
                    added += 1
        except BuildError as error:
            problems.append(f"Row {number}: {error}")
    return {"added": added, "updated": updated, "problems": problems}
